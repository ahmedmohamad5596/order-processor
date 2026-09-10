"""Seed the Egyptian areas layer (level 3: below governorate/city).

Writes the ''المناطق والأحياء'' sheet into egypt_governorates.xlsx with stable
IDs, parent relationships (governorate/city), entity types and aliases. Each
seed area is a well-known Egyptian geographic entity; learned areas from
areas_learned.json are merged on top at lookup-build time.

Run:  python -X utf8 scripts/seed_areas.py
"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl

from engine.lookup_builder import _EXCEL_PATH, build_lookup
from engine.normalizer import normalize_lookup_key

# (governorate, city, entity_type, area_name, aliases, confidence)
# city may be empty for governorate-level entries. Names are the exact
# display strings from the governorate/city tables.
AREA_SEED = [
    # ═══ القاهرة (01) ═══
    ("القاهرة", "القاهرة الجديدة", "area", "التجمع الأول", "التجمع 1", "high"),
    ("القاهرة", "القاهرة الجديدة", "area", "التجمع الثالث", "التجمع 3", "high"),
    ("القاهرة", "القاهرة الجديدة", "area", "التجمع الخامس", "التجمع 5", "high"),
    ("القاهرة", "القاهرة الجديدة", "area", "الرحاب", "", "high"),
    ("القاهرة", "القاهرة الجديدة", "area", "النرجس", "", "high"),
    ("القاهرة", "القاهرة الجديدة", "area", "القطامية", "", "high"),
    ("القاهرة", "القاهرة الجديدة", "area", "النزهة الجديدة", "", "medium"),
    ("القاهرة", "القاهرة الجديدة", "compound", "ريتاج", "", "medium"),
    ("القاهرة", "القاهرة الجديدة", "compound", "مدينتي", "", "medium"),
    ("القاهرة", "القاهرة الجديدة", "compound", "سراي القاهرة", "سراي", "medium"),
    ("القاهرة", "القاهرة الجديدة", "compound", "ماونتن فيو", "ماونتن فيو 3", "medium"),
    ("القاهرة", "القاهرة الجديدة", "compound", "فالي", "فاللي", "medium"),
    ("القاهرة", "القاهرة الجديدة", "compound", "وادي دجلة", "", "medium"),
    ("القاهرة", "شرق مدينة نصر", "area", "الحي الثامن", "", "high"),
    ("القاهرة", "شرق مدينة نصر", "area", "الحي السابع", "", "medium"),
    ("القاهرة", "شرق مدينة نصر", "area", "الحي العاشر", "", "medium"),
    ("القاهرة", "شرق مدينة نصر", "area", "مكرم عبيد", "", "medium"),
    ("القاهرة", "الزاوية الحمراء", "area", "بركة النصر", "", "high"),
    ("القاهرة", "الزيتون", "area", "حلمية الزيتون", "", "high"),
    ("القاهرة", "الشروق", "area", "الهايكستب", "هيكستب", "high"),
    ("القاهرة", "", "area", "مدينة السلام", "شرق مدينة السلام", "high"),
    ("القاهرة", "", "area", "منشأة ناصر", "", "high"),
    ("القاهرة", "", "area", "المطرية", "", "high"),
    ("القاهرة", "", "area", "شبرا", "شبرا القاهرة", "high"),
    ("القاهرة", "", "area", "العتبة", "العتبه", "medium"),
    ("القاهرة", "", "area", "باب اللوق", "", "medium"),
    ("القاهرة", "", "area", "الحسين", "", "medium"),
    ("القاهرة", "", "area", "الفسطاط", "", "medium"),
    ("القاهرة", "", "area", "العباسية", "", "high"),
    ("القاهرة", "", "area", "المعادي الجديدة", "", "medium"),
    ("القاهرة", "", "area", "المعادي القديمة", "", "medium"),
    ("القاهرة", "", "area", "حدائق المعادي", "", "medium"),
    ("القاهرة", "", "area", "عرب المعادي", "", "medium"),
    ("القاهرة", "", "area", "ميدان التحرير", "التحرير", "high"),
    ("القاهرة", "", "area", "ميدان رمسيس", "", "medium"),
    ("القاهرة", "", "area", "العاصمة الإدارية", "العاصمة", "medium"),
    # Roads — long corridors that must never be treated as areas
    ("القاهرة", "", "road", "جسر السويس", "", "high"),
    ("القاهرة", "", "road", "صلاح سالم", "", "high"),
    ("القاهرة", "", "road", "الطريق الدائري", "الدائري", "high"),
    ("القاهرة", "", "road", "كورنيش النيل", "", "high"),
    ("القاهرة", "", "road", "شارع رمسيس", "", "medium"),
    ("القاهرة", "", "road", "شارع الجمهورية", "", "medium"),
    ("القاهرة", "", "road", "شارع عبد العزيز", "", "medium"),
    ("القاهرة", "", "road", "شارع قصر النيل", "", "medium"),
    ("القاهرة", "", "road", "شارع 26 يوليو", "", "medium"),
    ("القاهرة", "", "road", "شارع عباس العقاد", "عباس العقاد", "medium"),
    ("القاهرة", "", "road", "شارع مصطفى النحاس", "مصطفى النحاس", "medium"),
    ("القاهرة", "", "road", "طريق النصر", "", "medium"),
    ("القاهرة", "", "road", "شارع العروبة", "", "medium"),
    ("القاهرة", "", "road", "كوبري أكتوبر", "", "medium"),
    ("القاهرة", "", "road", "محور 26 يوليو", "", "medium"),
    # ═══ الجيزة (02) ═══
    ("الجيزة", "الدقي", "area", "المهندسين", "", "high"),
    ("الجيزة", "العجوزة", "area", "كيت كات", "", "high"),
    ("الجيزة", "الطالبية", "area", "الزريبة", "", "high"),
    ("الجيزة", "أكتوبر (مدينة 6 أكتوبر)", "area", "الحي الثاني", "", "high"),
    ("الجيزة", "أكتوبر (مدينة 6 أكتوبر)", "area", "حدائق أكتوبر", "", "medium"),
    ("الجيزة", "أكتوبر (مدينة 6 أكتوبر)", "compound", "بالم هيلز أكتوبر", "بالم هيلز", "medium"),
    ("الجيزة", "", "area", "العمرانية", "", "high"),
    ("الجيزة", "", "area", "المنيب", "", "medium"),
    ("الجيزة", "", "area", "المعتمدية", "", "medium"),
    ("الجيزة", "", "area", "صفط اللبن", "", "medium"),
    ("الجيزة", "", "area", "ناهيا", "", "medium"),
    ("الجيزة", "", "area", "زنين", "", "medium"),
    ("الجيزة", "", "area", "ميدان الجيزة", "", "medium"),
    ("الجيزة", "الهرم", "road", "شارع الهرم", "", "high"),
    ("الجيزة", "فيصل", "road", "شارع فيصل", "", "high"),
    ("الجيزة", "الدقي", "road", "شارع جامعة الدول العربية", "", "medium"),
    ("الجيزة", "الدقي", "road", "شارع التحرير", "", "medium"),
    ("الجيزة", "", "road", "طريق الفيوم", "", "medium"),
    ("الجيزة", "", "road", "محور صفط اللبن", "", "medium"),
    # ═══ القليوبية (03) ═══
    ("القليوبية", "", "area", "بهتيم", "", "high"),
    ("القليوبية", "", "area", "الخصوص", "", "high"),
    ("القليوبية", "قليوب", "area", "القلج", "", "medium"),
    # ═══ الإسكندرية (04) ═══
    ("الإسكندرية", "", "area", "كفر عبده", "", "high"),
    ("الإسكندرية", "", "area", "رشدي", "", "high"),
    ("الإسكندرية", "", "area", "ستانلي", "", "high"),
    ("الإسكندرية", "", "area", "الحضرة", "", "high"),
    ("الإسكندرية", "", "area", "السيوف", "", "medium"),
    ("الإسكندرية", "", "area", "سيدي بشر", "", "high"),
    ("الإسكندرية", "", "area", "المكس", "", "medium"),
    ("الإسكندرية", "", "area", "أبيس", "", "high"),
    ("الإسكندرية", "", "area", "محطة الرمل", "", "high"),
    ("الإسكندرية", "", "area", "الشاطبي", "", "medium"),
    ("الإسكندرية", "", "area", "الإبراهيمية", "الابراهيميه", "medium"),
    ("الإسكندرية", "", "area", "فلمنج", "", "medium"),
    ("الإسكندرية", "", "area", "سبورتنج", "", "medium"),
    ("الإسكندرية", "", "area", "سان ستيفانو", "سنت ستيفانو", "medium"),
    ("الإسكندرية", "", "area", "زيزينيا", "", "medium"),
    ("الإسكندرية", "", "area", "مصطفى كامل", "", "medium"),
    ("الإسكندرية", "", "area", "باب شرق", "", "medium"),
    ("الإسكندرية", "", "area", "القباري", "", "medium"),
    ("الإسكندرية", "", "area", "العامرية", "", "medium"),
    ("الإسكندرية", "", "area", "رأس التين", "", "medium"),
    ("الإسكندرية", "", "area", "بحري", "بحرية", "medium"),
    ("الإسكندرية", "", "road", "كورنيش الإسكندرية", "", "high"),
    ("الإسكندرية", "", "road", "طريق المحمودية", "", "medium"),
    ("الإسكندرية", "", "road", "طريق الجيش", "", "medium"),
    # ═══ البحيرة (05) ═══
    ("البحيرة", "", "area", "نوبرية", "النوبارية", "medium"),
    # ═══ مطروح (06) ═══
    ("مطروح", "", "area", "سيدي عبد الرحمن", "", "high"),
    ("مطروح", "", "area", "الساحل الشمالي", "", "medium"),
    # ═══ دمياط (07) ═══
    ("دمياط", "", "area", "عزبة البرج", "", "high"),
    ("دمياط", "", "area", "شطا", "", "medium"),
    # ═══ الدقهلية (08) ═══
    ("الدقهلية", "المنصورة", "area", "سندوب", "", "high"),
    ("الدقهلية", "شربين", "village", "محلة أنجاق", "", "high"),
    ("الدقهلية", "بلقاس", "area", "جمصة", "", "medium"),
    # ═══ كفر الشيخ (09) ═══
    ("كفر الشيخ", "البرلس (بلطيم)", "area", "بلطيم", "", "high"),
    ("كفر الشيخ", "", "area", "مصيف بلطيم", "", "medium"),
    # ═══ الغربية (10) ═══
    ("الغربية", "المحلة الكبرى", "village", "محلة روح", "", "medium"),
    # ═══ المنوفية (11) ═══
    ("المنوفية", "قويسنا", "village", "ميت بره", "", "high"),
    ("المنوفية", "السادات", "area", "المنطقة الثالثة", "", "high"),
    # ═══ الشرقية (12) ═══
    ("الشرقية", "العاشر من رمضان", "area", "العاشر", "العاشر من رمضان", "high"),
    # ═══ البحر الأحمر (18) ═══
    ("البحر الأحمر", "مدينة الغردقة", "area", "سقالة", "", "high"),
    ("البحر الأحمر", "مدينة الغردقة", "area", "دهار", "", "high"),
    # ═══ الفيوم (20) ═══
    ("الفيوم", "", "area", "قارون", "", "medium"),
    # ═══ أسيوط (22) ═══
    ("أسيوط", "أسيوط", "village", "نزلة عبداللاه", "", "high"),
    ("أسيوط", "ديروط", "area", "ديروط الشريف", "", "medium"),
    # ═══ قنا (25) ═══
    ("قنا", "قنا", "area", "دندرة", "", "medium"),
    # ═══ الأقصر (26) ═══
    ("الأقصر", "الأقصر", "area", "الكرنك", "", "high"),
    # ═══ أسوان (27) ═══
    ("أسوان", "البسيلي / الرديسية / أبو سمبل وتوشكى (منطقة النوبة)", "area", "أبو سمبل", "", "high"),
    # ═══ طريق وطني (ممتد بين محافظات) ═══
    ("القاهرة", "", "road", "طريق مصر إسكندرية الزراعي", "الطريق الزراعي", "medium"),
    ("القاهرة", "", "road", "طريق مصر إسكندرية الصحراوي", "الطريق الصحراوي", "medium"),
    ("القاهرة", "", "road", "طريق القاهرة السويس", "", "medium"),
    ("القاهرة", "", "road", "طريق القاهرة الإسماعيلية", "", "medium"),
    ("القاهرة", "", "road", "طريق بلبيس", "بلبيس الصحراوي", "medium"),
    ("القاهرة", "", "road", "الطريق الزراعي الشرقي", "", "medium"),
    ("القاهرة", "", "road", "الطريق الصحراوي الشرقي", "", "medium"),
]

SHEET = "المناطق والأحياء"
COLUMNS = ["id", "entity_type", "governorate", "city", "area_name", "aliases", "confidence"]


def _resolve_parents():
    lookup = build_lookup()
    gov_by_norm = {normalize_lookup_key(g["name_ar"]): g["name_ar"] for g in lookup["governorates"].values()}
    cities_by_gov = {}
    for c in lookup["cities"].values():
        cities_by_gov.setdefault(c["governorate_id"], []).append(c)

    resolved = []
    unresolved = []
    for gov, city, etype, area_name, aliases, conf in AREA_SEED:
        gname = gov_by_norm.get(normalize_lookup_key(gov))
        if not gname:
            unresolved.append((gov, city, area_name))
            continue
        gid = next(g["id"] for g in lookup["governorates"].values() if g["name_ar"] == gname)
        cname = ""
        if city:
            for c in cities_by_gov.get(gid, []):
                if normalize_lookup_key(c["name_ar"]) == normalize_lookup_key(city):
                    cname = c["name_ar"]
                    break
            if not cname:
                unresolved.append((gov, city, area_name))
                continue
        resolved.append({
            "governorate": gname,
            "governorate_id": gid,
            "city": cname,
            "city_id": next((c["id"] for c in cities_by_gov.get(gid, []) if c["name_ar"] == cname), None),
            "entity_type": etype,
            "area_name": area_name,
            "aliases": aliases,
            "confidence": conf,
        })
    return resolved, unresolved


def write_sheet():
    resolved, unresolved = _resolve_parents()
    if unresolved:
        raise SystemExit(f"unresolved seed parents: {unresolved}")

    wb = openpyxl.load_workbook(_EXCEL_PATH)
    if SHEET in wb.sheetnames:
        del wb[SHEET]
    ws = wb.create_sheet(SHEET)

    ws.append(["المناطق والأحياء المصرية (Areas / Neighborhoods / Roads / Compounds / Villages)"])
    ws.append([
        "بذرة يدوية لمستوى المنطقة (تحت المحافظة والمدينة). كل كيان له معرّف ثابت ونوع كيان وعلاقة بأبيه. "
        "تُدمج مع مناطق areas_learned.json عند بناء اللوك أب — المكتسبة تكمل ولا تحذف النوع من البذرة."
    ])
    ws.append([])
    ws.append(COLUMNS)
    for i, entry in enumerate(resolved, start=1):
        ws.append([
            f"AR{i:04d}",
            entry["entity_type"],
            entry["governorate"],
            entry["city"],
            entry["area_name"],
            entry["aliases"],
            entry["confidence"],
        ])

    wb.save(_EXCEL_PATH)
    print(f"wrote {len(resolved)} seed areas to sheet '{SHEET}'")


if __name__ == "__main__":
    write_sheet()